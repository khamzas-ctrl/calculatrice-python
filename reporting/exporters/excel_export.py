from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


VERT = "00B050"
ROUGE = "FF0000"
BLEU_HEADER = "1F497D"
GRIS = "F2F2F2"


def sauvegarder_excel(rapport: dict, chemin: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ratios Prudentiels"

    # En-tête banque
    ws.merge_cells("A1:D1")
    ws["A1"] = f"{rapport['banque']} — {rapport['periode']}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=BLEU_HEADER)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "Date de rapport :"
    ws["B2"] = rapport["date_rapport"]

    # Tableau des ratios
    headers = ["Ratio", "Valeur", "Seuil réglementaire", "Statut"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLEU_HEADER)
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rapport["ratios"], 5):
        couleur = VERT if r["conforme"] else ROUGE
        statut = "CONFORME" if r["conforme"] else "NON CONFORME"
        ws.cell(row=i, column=1, value=r["nom"])
        ws.cell(row=i, column=2, value=f"{r['valeur'] * 100:.2f}%")
        ws.cell(row=i, column=3, value=f"{r['seuil'] * 100:.2f}%")
        cell_statut = ws.cell(row=i, column=4, value=statut)
        cell_statut.font = Font(bold=True, color=couleur)
        if i % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=GRIS)

    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 32

    wb.save(chemin)
